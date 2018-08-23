import os.path as opath
import os
import pandas as pd
from openpyxl import Workbook, load_workbook
import matplotlib.pyplot as plt
from matplotlib import gridspec
#
from __path_organizer import aggr_dpath, h1data_dpath, h1chart_dpath
from __common import HOURS, TERMINALS1, TERMINALS2
from __common import clists, mlists, FIGSIZE

FONT_SIZE = 13


def gen_xlsx_file(terminals, ps_count, ofpath):
    wb = Workbook()
    ws_count = wb.active
    ws_count.title = "Count"
    ori_cell, dest_cell = ws_count.cell(row=2, column=1), ws_count.cell(row=1, column=2)
    ori_cell.value = 'Origin'
    dest_cell.value = 'Destination'
    for i, tn in enumerate(terminals + ['Total']):
        o_cell, d_cell = ws_count.cell(row=i + 3, column=1), ws_count.cell(row=2, column=i + 2)
        for cell in [o_cell, d_cell]:
            cell.value = tn
    pl_sc, sl_sc, totalCount = {}, {}, 0
    for i, prevLoc in enumerate(terminals):
        sumCount = 0
        for j, startLoc in enumerate(terminals):
            k = prevLoc, startLoc
            count = ps_count[k] if k in ps_count else 0
            sumCount += count
            cell = ws_count.cell(row=i + 3, column=j + 2)
            cell.value = count
        totalCount += sumCount
        j += 1
        cell = ws_count.cell(row=i + 3, column=j + 2)
        cell.value = sumCount
        pl_sc[prevLoc] = sumCount
    for i, startLoc in enumerate(terminals):
        cell = ws_count.cell(row=8, column=i + 2)
        sumCount = 0
        for prevLoc in terminals:
            k = prevLoc, startLoc
            count = ps_count[k] if k in ps_count else 0
            sumCount += count
        cell.value = sumCount
        sl_sc[startLoc] = sumCount
    i += 1
    cell = ws_count.cell(row=8, column=i + 2)
    cell.value = totalCount
    #
    ws_ratio = wb.copy_worksheet(ws_count)
    ws_ratio.title = "Ratio"
    for i, prevLoc in enumerate(terminals):
        for j, startLoc in enumerate(terminals):
            cell = ws_ratio.cell(row=i + 3, column=j + 2)
            k = prevLoc, startLoc
            count = ps_count[k] if k in ps_count else 0
            cell.value = count / float(pl_sc[prevLoc])
    for i, tn in enumerate(terminals):
        cell = ws_ratio.cell(row=8, column=i + 2)
        cell.value = pl_sc[tn] / float(totalCount)
        #
        cell = ws_ratio.cell(row=i + 3, column=7)
        cell.value = sl_sc[tn] / float(totalCount)
    i += 1
    cell = ws_ratio.cell(row=i + 3, column=7)
    cell.value = totalCount/ float(totalCount)
    wb.save(filename=ofpath)


def arrange_datasets():
    yyyys = ['2009', '2010']
    for yyyy in yyyys:
        terminals = TERMINALS1 if yyyy != '2018' else TERMINALS2
        df = pd.read_csv(opath.join(aggr_dpath, 'apTrip-%s.csv' % yyyy))
        ydf = df.groupby(['previous_dropoff_loc', 'start_loc']).count()['taxi_id'].reset_index(name="count")
        ps_count = {}
        for prevLoc, startLoc, count in ydf.values:
            ps_count[prevLoc, startLoc] = count
        ofpath = opath.join(h1data_dpath, 'AirportFlow-%s.xlsx' % yyyy)
        gen_xlsx_file(terminals, ps_count, ofpath)
        #
        mdf = df.groupby(['month', 'previous_dropoff_loc', 'start_loc']).count()['taxi_id'].reset_index(name="count")
        mps_count = {}
        months = set()
        for month, prevLoc, startLoc, count in mdf.values:
            mps_count[month, prevLoc, startLoc] = count
            months.add(month)
        for m in months:
            ps_count = {}
            for prevLoc in terminals:
                for startLoc in terminals:
                    k = m, prevLoc, startLoc
                    ps_count[prevLoc, startLoc] = mps_count[k] if k in mps_count else 0
            ofpath = opath.join(h1data_dpath, 'AirportFlow-%sM%02d.xlsx' % (yyyy, m))
            gen_xlsx_file(terminals, ps_count, ofpath)
        #
        hdf = df.groupby(['hour', 'previous_dropoff_loc', 'start_loc']).count()['taxi_id'].reset_index(name="count")
        hps_count = {}
        hours = set()
        for hour, prevLoc, startLoc, count in hdf.values:
            hps_count[hour, prevLoc, startLoc] = count
            hours.add(hour)
        for h in hours:
            ps_count = {}
            for prevLoc in terminals:
                for startLoc in terminals:
                    k = h, prevLoc, startLoc
                    ps_count[prevLoc, startLoc] = hps_count[k] if k in hps_count else 0
            ofpath = opath.join(h1data_dpath, 'AirportFlow-%sH%02d.xlsx' % (yyyy, h))
            gen_xlsx_file(terminals, ps_count, ofpath)
        wb = Workbook()
        ofpath = opath.join(h1data_dpath, 'AirportFlow-HourRatio-%s.xlsx' % yyyy)
        for h in HOURS:
            ws = wb.create_sheet(title='H%02d' % h)
            ifpath = opath.join(h1data_dpath, 'AirportFlow-%sH%02d.xlsx' % (yyyy, h))
            wb1 = load_workbook(ifpath)            
            for row in wb1['Ratio'].rows:
                ws.append([cell.value for cell in row])
        wb.save(filename=ofpath)
        ofpath = opath.join(h1data_dpath, 'AirportFlow-HourCount-%s.xlsx' % yyyy)
        for h in HOURS:
            ws = wb.create_sheet(title='H%02d' % h)
            ifpath = opath.join(h1data_dpath, 'AirportFlow-%sH%02d.xlsx' % (yyyy, h))
            wb1 = load_workbook(ifpath)
            for row in wb1['Count'].rows:
                ws.append([cell.value for cell in row])
            os.remove(ifpath)
        wb.save(filename=ofpath)


def draw_chart():
    years = [2009, 2010]
    df = None
    for yyyy in years:
        fpath = opath.join(aggr_dpath, 'apTrip-%d.csv' % yyyy)
        df = pd.read_csv(fpath) if df is None else df.append(pd.read_csv(fpath))
    hdf = df.groupby(['year', 'hour', 'previous_dropoff_loc', 'start_loc']).count()['taxi_id'].reset_index(name="count")
    hours = sorted(list(set(hdf['hour'])))
    yhLocs_count = {}
    for year, hour, pt, st, count in hdf.values:
        yhLocs_count[year, hour, pt, st] = count
    hyLocs_ratio = {}
    for pt in TERMINALS1:
        for h in hours:
            #
            for y in years:
                sumTrips = sum(yhLocs_count[y, h, pt, st] for st in TERMINALS1)
                for st in TERMINALS1:
                    hyLocs_ratio[h, y, pt, st] = yhLocs_count[y, h, pt, st] / float(sumTrips) * 100
    #
    for pt in TERMINALS1:
        for st in TERMINALS1:
            img_ofpath = opath.join(h1chart_dpath, 'FlowChange-%s-%s.pdf' % (pt, st))
            plt.figure(figsize=FIGSIZE)
            gs = gridspec.GridSpec(2, 1, height_ratios=[3, 1])
            ax1 = plt.subplot(gs[0])
            ax1.set_ylabel('%', fontsize=FONT_SIZE)
            for i, y in enumerate([2009, 2010]):
                plt.plot(range(len(hours)), [hyLocs_ratio[h, y, pt, st] if (h, y, pt, st) in hyLocs_ratio else 0 for h in HOURS],
                         color=clists[i], marker=mlists[i])
            plt.legend(['2009', '2010'], ncol=1, fontsize=FONT_SIZE)
            plt.xticks(range(len(HOURS)), HOURS)
            ax1.tick_params(axis='both', which='major', labelsize=FONT_SIZE)
            plt.setp(ax1.get_xticklabels(), visible=False)
            #
            ax2 = plt.subplot(gs[1], sharex=ax1)
            ydata = []
            for h in HOURS:
                d2009 = hyLocs_ratio[h, 2009, pt, st] if (h, 2009, pt, st) in hyLocs_ratio else 0 
                d2010 = hyLocs_ratio[h, 2010, pt, st] if (h, 2010, pt, st) in hyLocs_ratio else 0 
                ydata.append(d2010 - d2009)
            plt.plot(range(len(hours)), ydata, color=clists[2], marker=mlists[2])
            plt.legend(['2010 - 2009'], ncol=1, fontsize=FONT_SIZE)
            ax2.tick_params(axis='both', which='major', labelsize=FONT_SIZE)
            ax2.set_ylabel('%', fontsize=FONT_SIZE)
            ax2.set_xlabel('Hour', fontsize=FONT_SIZE)
            plt.savefig(img_ofpath, bbox_inches='tight', pad_inches=0)



if __name__ == '__main__':
    # arrange_datasets()
    draw_chart()